#!/usr/bin/env python3
"""다른 도구로 만든 .pptx 의 글에 ppt-deck 문구 검사(lint.py)를 돌린다.

pptxgenjs 같은 별도 틀로 만든 덱도 문장 규칙(TONE · MARKUP · 제목 개수 예고 · 쉼표 …)은 똑같이 적용된다.
슬라이드마다 글 상자와 발표 노트를 뽑아 임시 아웃라인을 만들고 lint.py 를 부른다.
배치·구조 검사(첫 장 cover, 레이아웃 리듬, 제목 병렬 등)는 이 덱의 틀과 맞지 않으므로 빼고 보여 준다.
(2026-10-02, 소집면접 기출 덱 점검에서 손으로 하던 것을 스크립트로)

사용:  lint_pptx_text.py deck.pptx
"""
from __future__ import annotations
import os, re, subprocess, sys, tempfile

import yaml
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
# 아웃라인 구조에서 나오는 검사 — 남의 틀로 만든 덱에는 뜻이 없다
# GLYPH 는 ppt-deck 글꼴(Pretendard) 기준이라 다른 글꼴을 쓴 덱에서는 오탐이다(Apple SD Gothic Neo 의 ∠ 등)
SKIP = ("STRUCTURE", "PARALLEL", "LAYOUT", "RHYTHM", "VISUAL", "SOURCE", "DENSITY",
        "BULLET", "NOTES", "COUNT_SLIDES", "GLYPH")


def outline(path):
    slides = []
    prs = Presentation(path)
    for s in prs.slides:
        ts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = " ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
                if t.strip() and not t.strip().isdigit():        # 쪽번호는 뺀다
                    ts.append(t)
        # 제목 찾기 — 제목 개체 틀이 있으면 그것, 없으면 화면 위쪽 22% 안에서 가장 큰 글.
        # 예전처럼 장 전체에서 가장 큰 글을 제목으로 보면 데이터 장에서는 「42」「7.3」 같은 숫자가 뽑혀
        # 진짜 제목(「모델 내부는 볼 수 없었다」 같은 표어형)을 검사하지 못했다 (2026-10-27)
        title, best = "", 0
        ph = [sh for sh in s.shapes if sh.is_placeholder and sh.has_text_frame
              and sh.placeholder_format.type in (1, 3) and sh.text_frame.text.strip()]   # TITLE · CENTER_TITLE
        if ph:
            title = " ".join(p.text for p in ph[0].text_frame.paragraphs if p.text.strip())
        else:
            H = prs.slide_height or 1
            for top_only in (True, False):
                for sh in s.shapes:
                    if not sh.has_text_frame: continue
                    if top_only and (sh.top or 0) > 0.22 * H: continue
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            sz = r.font.size.pt if r.font.size else 0
                            if sz > best and r.text.strip() and not re.fullmatch(r"[\d.,%\s]+", r.text.strip()):
                                # 두 줄 제목(문단 둘)은 글상자 전체가 제목이다
                                best, title = sz, " ".join(q.text for q in sh.text_frame.paragraphs if q.text.strip())
                if title:
                    break
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        slides.append({"layout": "statement", "title": title,
                       "bullets": [t for t in ts if t != title], "notes": notes})
    return {"meta": {"title": os.path.basename(path), "palette": "forest", "density": "speaker"},
            "slides": slides}


def main(path):
    with tempfile.NamedTemporaryFile("w", suffix=".yaml", delete=False, encoding="utf-8") as f:
        yaml.safe_dump(outline(path), f, allow_unicode=True)
        tmp = f.name
    r = subprocess.run([sys.executable, os.path.join(HERE, "lint.py"), tmp],
                       capture_output=True, text=True)
    os.unlink(tmp)
    keep = [l for l in r.stdout.splitlines()
            if re.search(r"\[[A-Z_]+\]", l) and not any(f"[{c}]" in l for c in SKIP)]
    print(f"ppt-deck 문구 검사 — {os.path.basename(path)}")
    for l in keep: print(l)
    print(f"\n  경고 {len(keep)}건 (배치·구조 검사는 제외)")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
