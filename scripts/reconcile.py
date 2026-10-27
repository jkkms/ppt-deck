#!/usr/bin/env python3
"""배포본 대조 — 덮어쓰기 전에 사용자가 직접 고친 것이 있는지 확인한다.

원칙: 자료를 고치기 **전에 반드시** 확인한다. 세션 시작 때 한 번이 아니라
**쓰기 직전마다**. md5 는 빌드마다 달라지므로(타임스탬프) 판단 기준이 아니다 —
도형 단위 대조가 기준이다.

  reconcile.py dump    deck.pptx
  reconcile.py diff    배포본.pptx 내빌드.pptx
  reconcile.py deploy  내빌드.pptx 배포본.pptx [--first]

deploy 는 배포본이 마지막 배포 이후 바뀌었으면 **거부하고 차이를 보여 준다.**
그 차이가 사용자의 수정이다. 되돌리지 말고 생성 스크립트에 반영한 뒤 다시 배포한다.
"""
from __future__ import annotations
import argparse, hashlib, json, os, re, shutil, sys, zipfile

from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.join(HERE, "..", ".baselines")
P = lambda v: round(Emu(v).pt, 1)


def _key(path):
    return hashlib.sha1(os.path.abspath(path).encode()).hexdigest()[:16]


def dump(path):
    """도형 단위 지문. (쪽, x, y, w, h, 크기, 굵기, 글) — 위치·서식·문구를 모두 본다."""
    prs = Presentation(path)
    out = []
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            geo = (i, P(sh.left), P(sh.top), P(sh.width), P(sh.height))
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                prst = re.search(r'prstGeom prst="([^"]+)"', sh._element.xml)
                out.append(geo + (None, None, "<shape:%s>" % (prst.group(1) if prst else "?")))
                continue
            for para in sh.text_frame.paragraphs:
                # 같은 크기·굵기로 이어진 run 은 하나로 합친다. 파워포인트는 고쳐 저장할 때 맞춤법 표시 등으로
                # 한 문장을 run 여러 개로 쪼갠다(「답」「 」「분포로」…) — 글이 같아도 사용자 수정으로 잘못 잡혔다(2026-10-27)
                merged = []
                for r in para.runs:
                    k = (r.font.size.pt if r.font.size else None, bool(r.font.bold))
                    if merged and merged[-1][0] == k:
                        merged[-1][1] += r.text
                    else:
                        merged.append([k, r.text])
                for (sz, bold), text in merged:
                    out.append(geo + (sz, bold, text))
    return out


def _sortkey(row):
    """None(글 없는 도형의 크기)과 숫자가 섞인 줄도 정렬되게. 원래 sorted() 가 TypeError 로 멈췄다(2026-10-27)."""
    return tuple((0, v, "") if isinstance(v, (int, float)) and not isinstance(v, bool)
                 else (1, 0, str(v)) for v in row)


def dump_deep(path):
    """dump() 가 못 보는 것 — 채움색 · 글자색 · 세로 정렬 · 그림(내용 해시) · 발표 노트.
    2026-10-27: 사용자가 돌려준 발표 덱에서 표지 제목의 세로 정렬(위 → 가운데)만 바뀐 것을 dump() 로는
    못 잡았다(위치·크기·글이 그대로라서). 색만 바꾼 것, 그림만 바꿔 끼운 것, 노트만 고친 것도 같다.
    기준 덤프 형식은 그대로 두고(옛 기준과 비교가 깨지지 않게) 이 지문을 따로 둔다."""
    prs = Presentation(path)
    out = []
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            geo = (i, P(sh.left), P(sh.top), P(sh.width), P(sh.height))
            fill = ""
            try:
                if sh.fill.type == 1:
                    fill = str(sh.fill.fore_color.rgb)
            except Exception:
                pass
            anchor, cols = "", []
            if sh.has_text_frame:
                # 세로 정렬 + 문단 가로 정렬. 가로 정렬만 바꾼 편집(막대 위 값을 가운데로, 2026-10-27)도 잡는다
                algns = []
                for para in sh.text_frame.paragraphs:
                    pPr = para._p.pPr
                    al = (pPr.get("algn") if pPr is not None else None) or "l"
                    if para.text.strip() and (not algns or algns[-1] != al):
                        algns.append(al)
                anchor = (sh.text_frame._txBody.bodyPr.get("anchor") or "t") + ("|" + ",".join(algns) if algns else "")
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        try:
                            c = str(r.font.color.rgb)
                        except Exception:
                            c = ""
                        if not cols or cols[-1] != c:
                            cols.append(c)
            pic = ""
            if sh.shape_type == 13:
                pic = hashlib.sha1(sh.image.blob).hexdigest()[:12]
            out.append(geo + ("deep", fill, anchor, "/".join(cols), pic))
        notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ""
        out.append((i, 0, 0, 0, 0, "notes", "", "", "", notes))
    return out


def _deep_diff(a, b, label_a, label_b):
    da, db = set(map(tuple, dump_deep(a))), set(map(tuple, dump_deep(b)))
    only_a, only_b = sorted(da - db, key=_sortkey), sorted(db - da, key=_sortkey)
    if not only_a and not only_b:
        return 0
    print(f"  서식·그림·노트 차이 — {label_a} {len(only_a)}개 · {label_b} {len(only_b)}개")
    for tag, rows in ((label_a, only_a), (label_b, only_b)):
        for r in rows[:30]:
            i, x, y, w, h, kind, fill, anchor, cols, extra = r
            if kind == "notes":
                print(f"    {tag} s{i:02d} 노트 {extra[:60]!r}")
            else:
                print(f"    {tag} s{i:02d} ({x:6},{y:6}) {w:6}x{h:5} 채움 {fill or '-':6} "
                      f"정렬 {anchor or '-':7} 글자색 {cols or '-'} {('그림 ' + extra) if extra else ''}")
    print()
    return 1


def _edit_signals(path):
    """PowerPoint 가 남기는 편집 흔적. 결정적 증거는 아니지만 참고가 된다."""
    try:
        x = zipfile.ZipFile(path).read("docProps/app.xml").decode()
    except Exception:
        return {}
    g = lambda t: (re.search(rf"<{t}>(.*?)</{t}>", x) or [None, ""])[1]
    return {"Application": g("Application"), "TotalTime": g("TotalTime")}


def diff(a, b, label_a="배포본", label_b="내 빌드"):
    da, db = set(map(tuple, dump(a))), set(map(tuple, dump(b)))
    only_a, only_b = sorted(da - db, key=_sortkey), sorted(db - da, key=_sortkey)
    deep = _deep_diff(a, b, label_a, label_b)
    if not only_a and not only_b:
        print("  ✓ 도형 단위로 완전히 같다" if not deep else
              "  위치 · 크기 · 글은 같고 위의 서식 · 그림 · 노트만 다르다")
        return deep
    print(f"  {label_a}에만 있는 것 {len(only_a)}개 · {label_b}에만 있는 것 {len(only_b)}개\n")
    for tag, rows in ((f"{label_a} (사용자 수정일 가능성)", only_a), (f"{label_b}", only_b)):
        if not rows:
            continue
        print(f"  --- {tag}")
        for r in rows[:40]:
            i, x, y, w, h, sz, bold, t = r
            print(f"    s{i:02d} ({x:6},{y:6}) {w:6}x{h:5} "
                  f"{('%.1fpt' % sz) if sz else '   -  '} {'B' if bold else ' '} {t[:44]!r}")
        if len(rows) > 40:
            print(f"    … 외 {len(rows) - 40}개")
        print()
    return 1


def deploy(src, dest, first=False):
    os.makedirs(os.path.abspath(BASE_DIR), exist_ok=True)
    bpath = os.path.join(BASE_DIR, _key(dest) + ".json")
    if os.path.exists(dest):
        if not os.path.exists(bpath):
            if not first:
                raise SystemExit(
                    f"✗ 기준 덤프가 없다. 지금 덮어쓰면 사용자의 수정을 확인할 방법이 없다.\n"
                    f"  배포본을 먼저 검토하고, 내 빌드와 같다고 판단되면 --first 로 기준을 세워라.")
        else:
            base = [tuple(r) for r in json.load(open(bpath, encoding="utf-8"))]
            cur = [tuple(r) for r in dump(dest)]
            dpath = bpath[:-5] + ".deep.json"
            deep_changed = False
            if os.path.exists(dpath):      # 옛 기준에는 없다 — 있을 때만 본다
                dbase = {tuple(r) for r in json.load(open(dpath, encoding="utf-8"))}
                deep_changed = dbase != {tuple(r) for r in dump_deep(dest)}
            if set(base) != set(cur) or deep_changed:
                print(f"✗ 배포본이 마지막 배포 이후 바뀌었다 — 덮어쓰지 않았다.\n")
                only_cur = sorted(set(cur) - set(base), key=_sortkey)
                for r in only_cur[:40]:
                    i, x, y, w, h, sz, bold, t = r
                    print(f"    s{i:02d} ({x:6},{y:6}) {w:6}x{h:5} "
                          f"{('%.1fpt' % sz) if sz else '   -  '} {'B' if bold else ' '} {t[:44]!r}")
                if len(only_cur) > 40:
                    print(f"    … 외 {len(only_cur) - 40}개")
                sig = _edit_signals(dest)
                if sig:
                    print(f"\n  편집 흔적: {sig}")
                print("\n  이 차이가 사용자의 수정이다. 되돌리지 말고 아웃라인·빌더에 반영한 뒤")
                print("  다시 빌드해서 배포하라. 의도가 불분명하면 물어봐라.")
                return 1
    shutil.copyfile(src, dest)
    json.dump(dump(dest), open(bpath, "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(dump_deep(dest), open(bpath[:-5] + ".deep.json", "w", encoding="utf-8"), ensure_ascii=False)
    print(f"✓ 배포 완료 -> {dest}")
    print(f"  기준 덤프 갱신: {os.path.normpath(bpath)}")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("dump"); d.add_argument("pptx")
    f = sub.add_parser("diff"); f.add_argument("a"); f.add_argument("b")
    p_ = sub.add_parser("deploy"); p_.add_argument("src"); p_.add_argument("dest")
    p_.add_argument("--first", action="store_true")
    n = ap.parse_args()
    if n.cmd == "dump":
        for r in dump(n.pptx):
            print(r)
        sys.exit(0)
    sys.exit(diff(n.a, n.b) if n.cmd == "diff" else deploy(n.src, n.dest, n.first))
